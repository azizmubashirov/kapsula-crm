from django.http import HttpRequest
from django.http.response import HttpResponse as HttpResponse
from django.views.generic import TemplateView, UpdateView, DeleteView, DetailView
from web_project import TemplateLayout
from apps.product.models import (
    Product, 
    Meterial, 
    Color,
    ProductionProduct, 
    ProductColorQuantity,
    QuantityChange,
    QuantityChangeMeterial,
    CategoryMeterial
    )
from django.urls import reverse_lazy, reverse
from django.shortcuts import redirect
from django.template.response import TemplateResponse, HttpResponse
from django.db.models import Q
from django.core.paginator import Paginator
from django.db import transaction
from .forms import OrderForm, OrderClientForm, ProductionProductForm
from .models import Order, OrderCLient, OrderPayment, OrderYearFolder
from django.db import transaction
from apps.geo.models import Region, District
from decimal import Decimal
from apps.order.models import OrderCLient
import pandas as pd
from django.utils.timezone import now
from django.db.models import Sum, F, Value, DecimalField, ExpressionWrapper, Count
from django.db.models.functions import Coalesce
from datetime import datetime, date, timedelta
from django.views.decorators.csrf import csrf_exempt
from django.http import JsonResponse
from openpyxl import Workbook
from django.contrib import messages
import re

class OrderClientFolderView(TemplateView):
    form_class = OrderClientForm
    success_url = reverse_lazy('app-order-client-folder')
    
    def get_context_data(self, **kwargs):
        search_query = self.request.GET.get('q', '')
        region = int(self.request.GET.get("region", 0))
        district = int(self.request.GET.get("district", 0))
        product = int(self.request.GET.get("product", 0))
        
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        objs = OrderCLient.objects.order_by('name')
        if search_query:
            objs = objs.filter(
                Q(name__icontains=search_query)|
                Q(region__name_uz__icontains=search_query)|
                Q(district__name_uz__icontains=search_query)|
                Q(phone_number=search_query)
                                     )
        
        if product:
            objs = objs.filter(product__id=product)
        if region:
            objs = objs.filter(region_id=region)
            context['districts'] = District.objects.filter(region_id=region).order_by('name_uz')
        else:
            context['districts'] = []
            district = 0
        if district:
            objs = objs.filter(district_id=district)
    
        context['regions'] = Region.objects.order_by('name_uz')
        context['dis'] = district
        context['reg'] = region
        context['items'] = objs
        context['search_query'] = search_query
        context['form'] = self.form_class()
        context['products'] = Product.objects.order_by('ordering')
        context['product'] = product
        context['current_year'] = datetime.now().year
        return context
    
    def post(self, request, *args, **kwargs):
        form = self.form_class(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.save()
            form.save_m2m()
            return redirect(self.success_url)
        else:
            print("Main Form Errors:", form.errors)
        context = self.get_context_data()
        context['form'] = form
        return self.render_to_response(context)

def edit_folders(request, pk):
    year = request.GET.get("year")
    if request.POST:
        form = OrderClientForm(request.POST, instance=OrderCLient.objects.filter(pk=pk).first())
        if form.is_valid():
            form.save()
    return redirect(reverse_lazy('app-order-list') + f"?client={pk}" + f"&year={year}")

def delete_folders(request, pk):
    model = OrderCLient.objects.filter(pk=pk).first()
    model.delete()
    return redirect(reverse_lazy('app-order-client-folder'))

class OrderView(TemplateView):
    form_class = OrderForm
    success_url = reverse_lazy('app-order-list')
    
    def get_context_data(self, **kwargs):
        current_year = datetime.now().year
        client = int(self.request.GET.get("client", 0))
        date = self.request.GET.get("date")
        year = int(self.request.GET.get("year", current_year)) if int(self.request.GET.get("year")) else current_year
        if year and not OrderYearFolder.objects.filter(name=year, folder_id=client).first():
            OrderYearFolder.objects.create(
                name=current_year,
                folder_id=client
            )
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        if date:
            date = date.split(',')
            orders = Order.objects.filter(client_id=client, year=year, created_at__range=date).order_by('id')
        else:
            orders = Order.objects.filter(client_id=client, year=year).order_by('id')
        context['items'] = orders
        context['obj'] = OrderCLient.objects.filter(pk=client).first()
        context['form'] = self.form_class(instance=Order(client_id=client))
        context['date'] = date
        context['colors'] = Color.objects.order_by('name')
        context['folder_form'] = OrderClientForm(instance=context['obj'])
        context['years'] = range(current_year + 1, current_year + 10)  
        context['current_year'] = str(year)
        context['year_folders'] = OrderYearFolder.objects.filter(folder=context['obj'])
        return context
    
    def post(self, request, *args, **kwargs):
        current_year = datetime.now().year
        client = self.request.GET.get('client', 0)
        year = int(self.request.GET.get("year", current_year)) if int(self.request.GET.get("year")) else current_year
        form = self.form_class(request.POST, instance=Order(client_id=client, year=year))
        if form.is_valid():
            with transaction.atomic():
                model = form.save(commit=False)
                model.save()
                OrderPayment.objects.create(
                   order=model,
                   paid_price=request.POST.get('paid_amount') or 0.0,
                   paid_price_sum=request.POST.get('paid_amount_sum') or 0,
                   comment = request.POST.get('comment', '')
                )
            return redirect(self.success_url + f"?client={client}" + f"&year={year}")
        else:
            print("Main Form Errors:", form.errors)
        context = self.get_context_data()
        context['form'] = form
        return self.render_to_response(context)

def add_year(request, pk):
    if request.POST:
        data = request.POST
        model = OrderYearFolder.objects.create(
            name=data['year'],
            folder_id = pk
        )
    return redirect(reverse_lazy('app-order-list') + f"?client={pk}" + f"&year={model.name}")
                
def order_edit(request, pk):
    client = request.GET.get('client', 0)
    year = request.GET.get('year', 0)
    if request.POST:
        data = request.POST
        date = data['created_date'] + f"/{datetime.now().year}"
        with transaction.atomic():
            order = Order.objects.filter(pk=pk).first()
            if not data['product'] in [None, '', [], '0', 0]:
                order.product = data['product']
            if not data['color'] in [None, '', [], '0', 0]:
                order.color_obj_id = data['color']
            order.count = data['count']
            order.price = data['price']
            total_paid_price = 0
            for paid in request.POST.getlist('paid_price'):
                total_paid_price += Decimal(paid)
            order.total_price_order = Decimal(data['count']) * Decimal(data['price'])
            order.remainder = (Decimal(data['count']) * Decimal(data['price'])) - Decimal(total_paid_price)
            
            order.created_date = datetime.strptime(date, "%d/%m/%Y").date()
            order.save()
            
            OrderPayment.objects.filter(order_id = order.id).delete()
            x = 0
            for paid in request.POST.getlist('paid_price'):
                OrderPayment.objects.create(
                    order = order,
                    paid_price = paid,
                    comment = request.POST.getlist('paid_price-comment')[x]
                )
                x += 1
    return redirect(reverse_lazy('app-order-list') + f"?client={client}" + f"&year={year}")

class RemainsView(TemplateView):
    form_class = ProductionProductForm
    
    def custom_sort(self, product):
        number = ''.join([ch for ch in product.product.name if ch.isdigit()])
        return int(number)

    def get_context_data(self, **kwargs):
        search_query = self.request.GET.get('q', '')
        product = int(self.request.GET.get('product', 0))
        color = int(self.request.GET.get('color', 0))
        page_number = int(self.request.GET.get('page', 1))
        entries = int(self.request.GET.get("entries", 25))
        
        
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        products = ProductionProduct.objects.filter(product__is_active=True).order_by('product__ordering')
        
        # if search_query:
        #     products = products.filter(
        #         Q(product__name__icontains=search_query)|
        #         Q(product__color__name__icontains=search_query)
        #         )

        # if product > 0:
        #     products = products.filter(product_id=product)
            
        # paginator = Paginator(products, entries)
        context['items'] = products
        context['products'] = Product.objects.filter(is_active=True).order_by('ordering')
        context['meterials'] = CategoryMeterial.objects.all().order_by('name')
        context['colors'] = Color.objects.all()
        context['search_query'] = search_query
        context['color'] = color
        context['product'] = product
        context['entries'] = entries
        context['entries_list'] = [5,25,50,100,200]
        context['form'] = self.form_class()
        return context
    
class ProductHistoryView(TemplateView):
    
    def get_date_month(self):
        today = date.today()
        start_date = today.replace(day=1)

        if today.month == 12:
            next_month = date(today.year + 1, 1, 1)
        else:
            next_month = date(today.year, today.month + 1, 1)

        date_range = [
            start_date.strftime("%Y-%m-%d"),
            next_month.strftime("%Y-%m-%d")
        ]
        return date_range
    
    def get_context_data(self, **kwargs):
        today = now().date()
        product = int(self.request.GET.get('product', 0))

        date = self.request.GET.get("date")
        if date:
            date = date.split(',')
        else:
            date = self.get_date_month()
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        quantity_change = QuantityChange.objects.filter(
                change_type='subtract',
                created_at__range=date,
            ).values('user', 'created_at__date', 'region').annotate(
                total_quantity=Sum('quantity_changed'), 
                count_changes=Count('id')
            ).order_by('created_at__date', 'user', 'region')
       
        products = Product.objects.filter(is_active=True).order_by('ordering')
        context['products'] = products
        context['colors'] = Color.objects.order_by('name')
        context['products_filters'] = Product.objects.filter(is_active=True).order_by('ordering')

        context['items'] = quantity_change
        context['date'] = date
        context['pro'] = product
        return context
    
class StockView(TemplateView):
    
    def get_context_data(self, **kwargs):
        search_query = self.request.GET.get('q', '')
        page_number = int(self.request.GET.get('page', 1))
        entries = int(self.request.GET.get("entries", 25))
        
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        quantity_change_stocks = QuantityChangeMeterial.objects.all().order_by('-id')

        paginator = Paginator(quantity_change_stocks, entries)
        context['items'] = paginator.page(page_number)
        context['search_query'] = search_query
        context['entries'] = entries
        context['entries_list'] = [5,25,50,100,200]
        return context

class DebtorsView(TemplateView):
    
    def get_context_data(self, **kwargs):
        date = self.request.GET.get('date', datetime.now().date().strftime('%Y-%m-%d'))
        context = TemplateLayout.init(self, super().get_context_data(**kwargs))
        if date: 
            client_balances = Order.objects.filter(created_at__lt=date).values('client__id', 'client__name').annotate(
                total_remainder=Sum('remainder')
            )
        else:
            client_balances = Order.objects.values('client__id', 'client__name').annotate(
                total_remainder=Sum('remainder')
            )
        context['debtors'] = [client for client in client_balances if client['total_remainder'] > 0]
        context['creditors'] = [client for client in client_balances if client['total_remainder'] < 0]
        context['total_debt'] = sum(client['total_remainder'] for client in context['debtors'])
        context['total_credit'] = sum(client['total_remainder'] for client in context['creditors'])
        context['date'] = date
        return context

def add_production_product(request, pk):
    if request.POST:
        data = request.POST
        with transaction.atomic():
            product_color_quantity = ProductColorQuantity.objects.filter(pk=pk).first()
            product_color_quantity.quantity += int(data['quantity'])
            product_color_quantity.save()
            
            QuantityChange.objects.create(
                product=product_color_quantity.product.product,
                color = product_color_quantity.color,
                user = request.user,
                change_type = 'add',
                type = 'amount',
                quantity_changed = int(data['quantity'])
            )
            
            # QuantityChangeMeterial.objects.create(
            #     meterial_id = data['meterial'],
            #     user = request.user,
            #     change_type = 'subtract',
            #     type = 'amount',
            #     quantity_changed = int(data['meterial_amount'])
            # )

            # stock = Meterial.objects.filter(pk=data['meterial']).first()
            # stock.amount -= int(data['meterial_amount'])
            # stock.save()
            
            return redirect(reverse_lazy('app-production-product-list') + f'#section{product_color_quantity.product.product.id}')
    return redirect(reverse_lazy('app-production-product-list'))

def edit_production_product(request, pk):
    if request.POST:
        data = request.POST
        with transaction.atomic():
            product_color_quantity = ProductColorQuantity.objects.filter(pk=pk).first()
            product_color_quantity.quantity -= int(data['quantity'])
            product_color_quantity.save()
            
            QuantityChange.objects.create(
                product=product_color_quantity.product.product,
                color = product_color_quantity.color,
                user = data['user'],
                region_id = data['region'],
                change_type = 'subtract',
                type = 'amount',
                quantity_changed = int(data['quantity'])
            )
        return redirect(reverse_lazy('app-production-product-list') + f'#section{product_color_quantity.product.product.id}')
            
    return redirect(reverse_lazy('app-production-product-list'))

def change_price_production_product(request, pk):
    if request.POST:
        data = request.POST
        with transaction.atomic():
            product = Product.objects.filter(pk=pk).first()
            if product:
                product.pack = data.get('weight', 0)
                product.save()
    return redirect(reverse_lazy('app-production-product-list')  + f'#section{product.id}')

def delete_quantity_production_product(request, pk):
    with transaction.atomic():
        product_color_quantity = ProductColorQuantity.objects.filter(pk=pk).first()
        product_color_quantity.quantity = 0
        product_color_quantity.save()
    return redirect(reverse_lazy('app-production-product-list') + f"#section{product_color_quantity.product.product.id}")

def order_payment(request, pk):
    client = request.GET.get('client')
    year = request.GET.get('year')
    if request.POST:
        data = request.POST
        with transaction.atomic():
            order = Order.objects.filter(pk=pk).first()
            if order and data['inlineRadioOptions'] == 'plus':
                OrderPayment.objects.create(
                    paid_price = data['amount'],
                    order = order,
                    comment = data['comment']
                )
                order.remainder = order.remainder - Decimal(data['amount'])
            elif order and data['inlineRadioOptions'] == 'minus':
                order_payment = OrderPayment.objects.filter(order=order).last()
                if order_payment.paid_price == Decimal(data['amount']):
                    if  len(OrderPayment.objects.filter(order=order)) == 1:
                        order_payment.paid_price = 0.0
                        order_payment.save()
                    else:
                        order_payment.delete()
                else:
                    order_payment.paid_price -= Decimal(data['amount'])
                    order_payment.save()
                order.remainder = order.remainder + Decimal(data['amount'])
            order.save()
    return redirect(reverse_lazy('app-order-list') + f"?client={client}" + f"&year={year}")

def add_payment(request):
    client = request.GET.get('client')
    year = request.GET.get('year')
    if request.POST:
        data = request.POST
        with transaction.atomic():
            order = Order.objects.create(
                client_id = client,
                remainder = -Decimal(data['amount']) if data['inlineRadioOptions'] == 'plus' else Decimal(data['amount']),
                )
            OrderPayment.objects.create(
                paid_price = data['amount'],
                order = order,
                comment = data['comment']
            )
    return redirect(reverse_lazy('app-order-list') + f"?client={client}" + f"&year={year}")

@csrf_exempt
def edit_shipment(request):
    if request.method == "POST":
        shipment_id = request.POST.getlist("shipment_id")
        user = request.POST.get("client")
        quantities = request.POST.getlist("count")
        product_ids = request.POST.getlist("product")
        color_ids = request.POST.getlist("color")

        stocks = QuantityChange.objects.filter(id__in=shipment_id)

        for index, stock in enumerate(stocks):
            old_quantity = stock.quantity_changed 
            old_color_id = stock.color_id 
            old_product_id = stock.product_id

            stock.user = user 
            stock.product_id = product_ids[index]
            stock.color_id = color_ids[index]

            if not quantities[index] or quantities[index] in ['', None, '0']:
                new_quantity = 0
            else:
                new_quantity = int(quantities[index])

            stock.quantity_changed = new_quantity 
            stock.save()

            new_color_id = stock.color_id 
            new_product_id = stock.product_id 

            if old_color_id and old_color_id != new_color_id:
                old_production_product = ProductColorQuantity.objects.filter(
                    product__product_id=old_product_id,
                    color_id=old_color_id
                ).first()

                if old_production_product:
                    old_production_product.quantity += old_quantity 
                    old_production_product.save()

            if new_color_id:
                new_production_product = ProductColorQuantity.objects.filter(
                    product__product_id=new_product_id,
                    color_id=new_color_id
                ).first()

                if new_production_product:
                    new_production_product.quantity -= new_quantity
                    new_production_product.save()

        return JsonResponse({"status": "success", "message": "Shipment updated!"})

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)
        

def export_order_to_excel(request):
    client = int(request.GET.get('client', 0))
    orders = Order.objects.filter(client_id=client).order_by('id')

    data = []
    
    for order in orders:
        data.append({
            'Дата': order.created_at.strftime('%d-%b/%Y'),
            'Продукт': order.product,
            'Количество': order.count,
            'Цена': order.price,
            'Сумма': order.count * order.price,
            'Оплата (сум)': order.order_payments.all()[0].paid_price_sum,
            'Оплата': order.order_payments.all()[0].paid_price,
            '': '',
            'Остаток': order.remainder,
            'Комментарий': ''
        })
        for obj in order.order_payments.all()[1:]:
            data.append({
                'Дата': '',
                'Продукт': '',
                'Количество': '',
                'Цена': '',
                'Сумма': '',
                'Оплата (сум)': obj.paid_price_sum,
                'Оплата': obj.paid_price,
                '': obj.created_at.strftime('%d-%b/%Y'),
                'Остаток': order.remainder,
                'Комментарий': obj.comment or ''
            })
            
    data.append({
                'Дата': '',
                'Продукт': '',
                'Количество': '',
                'Цена': '',
                'Сумма': '',
                'Оплата (сум)': '',
                'Оплата': '',
                '': '',
                'Остаток': calculation_remainder(orders),
                'Комментарий': ''
            })
    df = pd.DataFrame(data)

    file_path = "media/analytics.xlsx"
    with pd.ExcelWriter(file_path) as writer:
        df.to_excel(writer, sheet_name='Распродажа', index=False)

    with open(file_path, "rb") as excel_file:
        response = HttpResponse(excel_file.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=order.xlsx"
        return response
    
def export_debtors_to_excel(request):
    client_balances = Order.objects.values('client__id', 'client__name').annotate(
        total_remainder=Sum('remainder')
    )
    debtors = [client for client in client_balances if client['total_remainder'] > 0]
    creditors = [client for client in client_balances if client['total_remainder'] < 0]

    data = []

    max_len = max(len(debtors), len(creditors))  # Eng uzun ro'yxat uzunligini olish

    for i in range(max_len):
        row = {
            'Qarzdorlar': debtors[i]['client__name'] if i < len(debtors) else '',
            'Qarzdor Summasi': debtors[i]['total_remainder'] if i < len(debtors) else '',
            'Xaqdorlar': creditors[i]['client__name'] if i < len(creditors) else '',
            'Xaqdor Summasi': creditors[i]['total_remainder'] if i < len(creditors) else '',
        }
        data.append(row)
    df = pd.DataFrame(data)

    file_path = "media/analytics.xlsx"
    with pd.ExcelWriter(file_path) as writer:
        df.to_excel(writer, sheet_name='document', index=False)
        # worksheet = writer.sheets["document"]
        # worksheet.set_column("A:A", 20)  
        # worksheet.set_column("B:B", 15) 
        # worksheet.set_column("C:C", 20) 
        # worksheet.set_column("D:D", 15)

    with open(file_path, "rb") as excel_file:
        response = HttpResponse(excel_file.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=document.xlsx"
        return response

def import_clients_from_excel(request):
    if request.method == 'POST' and request.FILES.get('myfile'):
        myfile = request.FILES['myfile']
        try:
            df = pd.read_excel(myfile, engine="openpyxl")

            df = df.dropna(how='all')
            
            for _, row in df.iterrows():
                name = row.iloc[1] 
                comment = row.iloc[6] if len(row) > 6 else None 
                phone_number = row.iloc[3].split()[0] if pd.notna(row.iloc[3]) else None 

                if pd.notna(name): 
                    OrderCLient.objects.create(
                        name=name,
                        comment=comment,
                        phone_number=phone_number,
                    )
            messages.success(request, "Mijozlar muvaffaqiyatli yuklandi!")
        except Exception as e:
            messages.error(request, "Faylni yuklashda muammo bor! Boshqa fayl yuklang", 'danger')

        return redirect(reverse_lazy('app-order-client-folder'))
    
def calculation_remainder(orders):
    total = 0
    for obj in orders:
        total += obj.remainder
    return total

def delete_order(request, pk):
    client = request.GET.get('client', 0)
    year = request.GET.get('year', 0)
    order = Order.objects.get(pk=pk)
    order.delete()
    return redirect(reverse_lazy('app-order-list') + f"?client={client}" + f"&year={year}")

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import datetime
from apps.product.models import Product
from django.db.models import Sum

def get_product_history_items(date_range, product_id=None):
    return QuantityChange.objects.filter(
                change_type='subtract',
                created_at__range=date_range,
            ).values('user', 'created_at__date', 'region').annotate(
                total_quantity=Sum('quantity_changed'), 
                count_changes=Count('id')
            ).order_by('created_at__date', 'user', 'region')

def get_product_quantity(item, product, color):
    return QuantityChange.objects.filter(
        color=color, product=product,
        change_type='subtract',
        created_at__date=item['created_at__date'],
        user = item['user'],
        region=item['region']
    ).aggregate(total_quantity=Sum('quantity_changed'))['total_quantity'] or 0

def total_quantity_changes(date_range, product_id=None):
    items = get_product_history_items(date_range, product_id)
    products = Product.objects.all().prefetch_related('color')
    totals = []
    
    for product in products:
        for color in product.color.all():
            total = sum(get_product_quantity(item, product, color) for item in items)
            totals.append(total)
        totals.append('')
    
    return totals

def total_product_quantity_changes(date_range, product_id=None):
    items = get_product_history_items(date_range, product_id)
    products = Product.objects.all().prefetch_related('color')
    totals = []
    
    for product in products:
        product_total = 0
        for color in product.color.all():
            color_total = sum(get_product_quantity(item, product, color) for item in items)
            totals.append(color_total)
            product_total += color_total
        totals.append(product_total)
    
    return totals

def total_kg_changes(date_range, product_id=None):
    items = get_product_history_items(date_range, product_id)
    products = Product.objects.all().prefetch_related('color')
    totals = []
    
    for product in products:
        for color in product.color.all():
            total = sum(get_product_quantity(item, product, color) * product.weight for item in items)
            totals.append(total)
        totals.append('')
    
    return totals

def total_product_kg_changes(date_range, product_id=None):
    items = get_product_history_items(date_range, product_id)
    products = Product.objects.all().prefetch_related('color')
    totals = []
    
    for product in products:
        product_total = 0
        for color in product.color.all():
            color_total = sum(get_product_quantity(item, product, color) * product.weight for item in items)
            totals.append(color_total)
            product_total += color_total
        totals.append(product_total)
    
    return totals

def export_product_history(request):
    date = request.GET.get('date', '')
    pro = request.GET.get('pro')
    today = now().date()
    if not date:
        date = [f"{today.year}-{today.month:02d}-01", f"{today.year}-{today.month:02d}-28"]
    else:
        date = date.split(',')
        
    wb = Workbook()
    ws = wb.active
    ws.title = 'Product History'
    
    products = Product.objects.all().prefetch_related('color')
    items = get_product_history_items(date, pro)
    
    ws.column_dimensions['A'].width = 15 
    ws.column_dimensions['B'].width = 30
    
    current_col = 3 
    headers = [['время', 'Организация', '']]
    sub_headers = [['', '', '']]
    
    ws.cell(row=1, column=1, value='время')
    ws.cell(row=1, column=2, value='Организация')
    
    for product in products:
        colors_count = product.color.count()
        ws.cell(row=1, column=current_col, value=product.name)
        if colors_count > 1:
            ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + colors_count - 1)
        
        for color in product.color.all():
            cell = ws.cell(row=2, column=current_col)
            cell.value = color.name
            if color.code:
                cell.fill = PatternFill(start_color=color.code.replace('#', ''), end_color=color.code.replace('#', ''), fill_type='solid')
                cell.font = Font(color='FFFFFF' if color.code == '#000000' else '000000')
            current_col += 1
        
        ws.cell(row=1, column=current_col, value='')
        ws.cell(row=2, column=current_col, value='')
        current_col += 1
    
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal='center')
    
    current_row = 3
    for item in items:
        ws.cell(row=current_row, column=1, value=item['created_at__date'].strftime('%d-%b'))
        ws.cell(row=current_row, column=2, value=str(item['user']))
        
        current_col = 3
        for product in products:
            for color in product.color.all():
                quantity = get_product_quantity(item, product, color)
                ws.cell(row=current_row, column=current_col, value=quantity)
                current_col += 1
            current_col += 1
        
        current_row += 1
    
    ws.cell(row=current_row, column=1, value='')
    ws.cell(row=current_row, column=2, value='')
        
    current_col = 3
    for product in products:
        for obj in product.color.all():
            count = QuantityChange.objects.filter(
                color=obj, product=product,
                change_type='subtract',
                created_at__range=date
            ).aggregate(total=Sum('quantity_changed'))['total'] or 0
            ws.cell(row=current_row, column=current_col, value=count)
            current_col += 1
        current_col += 1
    current_row += 1

    current_col = 3
    for product in products:
        count = QuantityChange.objects.filter(
            product=product,
            change_type='subtract',
            created_at__range=date
        ).aggregate(total=Sum('quantity_changed'))['total'] or 0
        colors_count = product.color.count()
        if colors_count > 0:
            ws.merge_cells(start_row=current_row, start_column=current_col,
                          end_row=current_row, end_column=current_col + colors_count - 1)
        ws.cell(row=current_row, column=current_col, value=f"{count} шт")
        current_col += colors_count + 1 if colors_count > 0 else 1
    current_row += 1

    current_col = 3
    for product in products:
        for obj in product.color.all():
            count = QuantityChange.objects.filter(
                color=obj, product=product,
                change_type='subtract',
                created_at__range=date
            ).aggregate(total=Sum('quantity_changed'))['total'] or 0
            number_part = re.match(r'^\d+', product.name)
            if number_part:
                number_part = number_part.group()
            try:
                result = int(number_part) / 1000 if number_part else 0.01
            except ValueError:
                result = 0.01
            total_weight = count * result
            ws.cell(row=current_row, column=current_col, value=total_weight)
            current_col += 1
        current_col += 1
    current_row += 1

    current_col = 3
    for product in products:
        count = QuantityChange.objects.filter(
            product=product,
            change_type='subtract',
            created_at__range=date
        ).aggregate(total=Sum('quantity_changed'))['total'] or 0
        number_part = ''.join([char for char in product.name if char.isdigit()])[:2]
        try:
            result = int(number_part) / 1000 if number_part else 0.01
        except ValueError:
            result = 0.01
        weight = product.weight if product.weight else result
        total_weight = count * weight
        colors_count = product.color.count()
        if colors_count > 0:
            ws.merge_cells(start_row=current_row, start_column=current_col,
                          end_row=current_row, end_column=current_col + colors_count - 1)
            ws.cell(row=current_row, column=current_col, value=f"{round(total_weight, 2)} кг")
        else:
            ws.cell(row=current_row, column=current_col, value=f"{round(total_weight, 2)} кг")
        current_col += colors_count + 1 if colors_count > 0 else 1
    current_row += 1    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))
    for row in ws.iter_rows(min_row=1, max_row=current_row-1):
        for cell in row:
            cell.border = thin_border
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=otgruska_{}.xlsx'.format(
        datetime.now().strftime('%Y%m%d_%H%M%S')
    )
    
    wb.save(response)
    return response


def export_to_excel_production_product(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mahsulotlar"

    products = ProductionProduct.objects.prefetch_related('product_color_quantity__color')

    row_idx = 1
    
    for product in products:
        ws.cell(row=row_idx, column=1, value=product.product.name).font = Font(bold=True)
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")

        for obj in product.product_color_quantity.all():
            color_name = obj.color.name
            quantity = obj.quantity
            color_hex = obj.color.code.lstrip('#')

            ws.cell(row=row_idx, column=2, value=color_name)
            ws.cell(row=row_idx, column=3, value=quantity).alignment = Alignment(horizontal="center")

            fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
            ws.cell(row=row_idx, column=4, value=obj.color.code).fill = fill

            row_idx += 1 

        row_idx += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=ostatka_{}.xlsx'.format(
        datetime.now().strftime('%Y%m%d_%H%M%S')
    )
    wb.save(response)

    return response